#!/usr/bin/env python3
"""Apply frozen source locality decisions and retain all unresolved institutions."""
import argparse, csv, json, hashlib, re, unicodedata
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'workspace/reports/geography-research-restoration/2026-10-13'

def norm(value):
    return re.sub(r'[^\w]+', ' ', ''.join(c for c in unicodedata.normalize('NFKD', str(value or '')).casefold() if not unicodedata.combining(c))).strip()

def records(name):
    return [json.loads(line) for line in (REPORT / name).read_text().splitlines() if line]

def save_records(name, rows):
    (REPORT / name).write_text(''.join(json.dumps(row, ensure_ascii=False) + '\n' for row in rows))

def apply(additional_decisions=None):
    summary = json.loads((REPORT / 'summary.json').read_text())
    if summary.get('restorationsApplied') and additional_decisions is None:
        raise ValueError('Restorations already applied')
    previous = json.loads((REPORT / 'proposed-university-restorations.json').read_text())
    selected = json.loads(Path(additional_decisions).read_text()) if additional_decisions else previous
    decisions = {row['referenceId']: row for row in selected}
    if len(decisions) != len(selected):
        raise ValueError('Duplicate decision identity')
    if additional_decisions and set(decisions) & {row['referenceId'] for row in previous}:
        raise ValueError('Already restored university decision')
    for artifact in summary['artifacts']:
        if hashlib.sha256((ROOT / artifact['path']).read_bytes()).hexdigest() != artifact['afterSha256']:
            raise ValueError('Source changed after the reviewed snapshot')
    region_rows = [row for artifact in summary['artifacts'] if artifact['kind'] == 'regions' for row in csv.DictReader((ROOT / artifact['path']).open(encoding='utf-8-sig'))]
    regions = {row['regionCode']: row for row in region_rows}
    city_files = {artifact['path']: list(csv.DictReader((ROOT / artifact['path']).open(encoding='utf-8-sig'))) for artifact in summary['artifacts'] if artifact['kind'] == 'cities'}
    cities = {row['cityId']: row for rows in city_files.values() for row in rows}
    city_changes = records('city-changes.jsonl')
    university_changes = records('university-changes.jsonl')
    def bind_city(decision):
        place = decision['place']
        if regions[place['regionCode']]['countryIso3'] != decision['countryIso3']:
            raise ValueError('Cross-country source region decision')
        if not decision['source'].get('url'):
            raise ValueError('Missing decision evidence')
        if place.get('existingCityId'):
            city = cities[place['existingCityId']]
            if city['countryIso3'] != decision['countryIso3']:
                raise ValueError('Cross-country source decision')
            return city
        name, code = decision['name'], place['regionCode']
        candidates = [city for city in cities.values() if city['countryIso3'] == decision['countryIso3'] and city['regionCode'] == code and norm(name) in {norm(city[k]) for k in ('cityNameEn', 'cityAscii', 'localName')} | {norm(a) for a in city['cityAliases'].split('|') if a}]
        if len(candidates) == 1:
            return candidates[0]
        if len(candidates) > 1:
            if place.get('id'):
                exact = [city for city in candidates if city['sourceCityId'] == 'geonames:' + place['id']]
                if len(exact) == 1:
                    return exact[0]
            raise ValueError('Ambiguous existing city identity: ' + decision['referenceId'] + ' / ' + name + ' / ' + ','.join(c['cityId'] for c in candidates))
        if place.get('id'):
            exact = [city for city in cities.values() if city['countryIso3'] == decision['countryIso3'] and city['sourceCityId'] == 'geonames:' + place['id']]
            if len(exact) == 1 and exact[0]['regionCode'] == code:
                return exact[0]
        region = regions[code]
        filename = next(path for path, rows in city_files.items() if any(row['countryIso2'] == region['countryIso2'] for row in rows))
        template = city_files[filename][0]
        city = {key: '' for key in template}
        city.update({key: region[key] for key in ('countryIso2', 'countryIso3', 'countryNameEn', 'countryNameAr', 'continent', 'continentAr', 'subregion', 'subregionAr')})
        identity = 'city_' + region['countryIso2'].lower() + ('_geonames_' + place['id'] if place.get('id') else '_official_' + norm(name).replace(' ', '_'))
        if identity in cities:
            raise ValueError('Conflicting new city identity')
        city.update({'importAction': 'UPSERT', 'cityId': identity, 'sourceCityId': 'geonames:' + place['id'] if place.get('id') else 'official:' + region['countryIso2'] + ':' + norm(name).replace(' ', '_'),
                     'cityNameEn': name, 'cityAscii': place.get('ascii', name), 'localName': place.get('name', name),
                     'regionCode': code, 'regionNameEn': region['nameEn'], 'regionNameAr': region['nameAr'], 'regionMatchStatus': 'MATCHED_BY_CODE',
                     'cityType': 'UNIVERSITY_LOCALITY', 'isCountryCapital': 'false', 'isAdministrativeCapital': 'false', 'isMajorCity': 'false',
                     'isStudyCityCandidate': 'true', 'isUniversityCityCandidate': 'true', 'isTestCenterCityCandidate': 'NEEDS_REVIEW',
                     'latitude': place.get('latitude', ''), 'longitude': place.get('longitude', ''), 'timezone': place.get('timezone', ''), 'population': place.get('population', ''),
                     'sourceType': 'MANUAL_FILE_SOURCE', 'sourceUrl': decision['source']['url'], 'sourceDataset': 'Reviewed locality source decisions, 2026-10-13; GeoNames or primary-address evidence',
                     'verificationStatus': 'NEEDS_MANUAL_REVIEW', 'translationStatus': 'NEEDS_AR_TRANSLATION', 'notes': 'Source locality: ' + decision['scope'] + '. No canonical FK or publication approval. See geography-research-restoration/2026-10-13 university receipts.'})
        city_files[filename].append(city)
        cities[identity] = city
        city_changes.append({'action': 'ADD_CITY', 'cityId': identity, 'after': dict(city), 'sourceUrl': decision['source']['url'], 'gazetteer': place if place.get('id') else None, 'universityReferenceId': decision['referenceId']})
        return city
    remaining = []
    accepted = 0
    prepared_workbooks = []
    applied_refs = set()
    for artifact in summary['artifacts']:
        if artifact['kind'] != 'universities':
            continue
        path = ROOT / artifact['path']
        workbook = openpyxl.load_workbook(path)
        active, review = workbook.worksheets[0], workbook['Geography Review']
        active_header_index = next(i for i, values in enumerate(active.values, 1) if 'Reference ID' in values)
        active_headers = [cell.value for cell in active[active_header_index]]
        review_header_index = next(i for i, values in enumerate(review.values, 1) if 'Reference ID' in values)
        review_headers = [cell.value for cell in review[review_header_index]]
        archive_path = ROOT / 'workspace/reports/active-geography-alignment/2026-10-13' / path.name
        archive = openpyxl.load_workbook(archive_path, read_only=True, data_only=True) if not additional_decisions else None
        archive_sheet = archive['Geography Review'] if archive else review
        source_rows = list(archive_sheet.values)
        if additional_decisions:
            review.delete_rows(review_header_index + 1, review.max_row - review_header_index)
        headers = None
        for values in source_rows:
            if 'Reference ID' in values:
                headers = list(values)
                continue
            if not headers:
                continue
            original = dict(zip(headers, values))
            ref = original.get('Reference ID')
            if not ref:
                continue
            row = dict(original)
            decision = decisions.get(ref)
            if decision:
                city = bind_city(decision)
                row.update({'City': city['cityNameEn'], 'City Source ID': city['cityId'], 'City Region Code': city['regionCode'], 'Geography Status': 'SOURCE_MATCHED',
                            'Geography Review Reason': 'REGISTRY_CITY_CORROBORATED_MAIN_CAMPUS_NOT_LIVE_VERIFIED' if decision['method'] == 'USER_CITY_AND_ROR_REGISTRY_CORROBORATED' else None})
                active.append([row.get(header) for header in active_headers])
                university_changes.append({'referenceId': ref, 'action': 'RESTORE_CORROBORATED_LOCALITY', 'before': original, 'after': row, 'afterSheet': active.title,
                                           'citySourceId': city['cityId'], 'method': decision['method'], 'scope': decision['scope'], 'source': decision['source'], 'addressExcerpt': decision.get('addressExcerpt'),
                                           'registryRecord': decision.get('registryRecord')})
                accepted += 1
                applied_refs.add(ref)
            else:
                review.append([row.get(header) for header in review_headers])
                remaining.append({'referenceId': ref, 'name': original['Official English Name'], 'countryIso3': original['ISO3'], 'city': original['City'], 'reason': original['Geography Review Reason'], 'sourcePath': artifact['path']})
        if archive:
            archive.close()
        prepared_workbooks.append((path, workbook))
    if applied_refs != set(decisions):
        raise ValueError('Decision does not identify an unresolved university')
    # Validate every decision before writing any source workbook.
    for path, workbook in prepared_workbooks:
        workbook.save(path)
        workbook.close()
    for filename, rows in city_files.items():
        with (ROOT / filename).open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator='\n')
            writer.writeheader()
            writer.writerows(rows)
    for artifact in summary['artifacts']:
        path = ROOT / artifact['path']
        artifact['afterSha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        if artifact['kind'] == 'cities':
            artifact['afterRows'] = len(city_files[artifact['path']])
    summary.update({'restorationsApplied': True, 'restoredUniversities': accepted + (summary['restoredUniversities'] if additional_decisions else 0), 'remainingReviewUniversities': len(remaining), 'originalUniversityRows': 10723})
    (REPORT / 'summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2) + '\n')
    save_records('city-changes.jsonl', city_changes)
    save_records('university-changes.jsonl', university_changes)
    save_records('remaining-university-review.jsonl', remaining)
    if additional_decisions:
        (REPORT / 'proposed-university-restorations.json').write_text(json.dumps(previous + selected, ensure_ascii=False, indent=2) + '\n')
    with (REPORT / 'remaining-university-review.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=['referenceId', 'name', 'countryIso3', 'city', 'reason', 'sourcePath'], lineterminator='\n')
        writer.writeheader()
        writer.writerows(remaining)
    print(json.dumps({'restoredUniversities': accepted, 'remainingReview': len(remaining), 'activeCities': len(cities), 'databaseWrites': 0}))

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--additional-decisions', type=Path)
    apply(parser.parse_args().additional_decisions)

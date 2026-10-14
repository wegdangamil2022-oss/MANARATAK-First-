import importlib.util,json,hashlib,shutil,tempfile,unittest,csv,io,subprocess
from pathlib import Path
import openpyxl
script=Path(__file__).resolve().parents[1]/'verify-university-city-research.py';spec=importlib.util.spec_from_file_location('city_research',script);checker=importlib.util.module_from_spec(spec);spec.loader.exec_module(checker);originalRoot,originalReport=checker.ROOT,checker.REPORT
class ResearchIntegrityTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();checker.ROOT=Path(self.tmp.name);checker.REPORT=checker.ROOT/'workspace/reports/university-city-research/2026-10-13';(checker.ROOT/'.git').symlink_to(originalRoot/'.git',target_is_directory=True)
  shutil.copytree(originalRoot/'workspace/reference-data',checker.ROOT/'workspace/reference-data');shutil.copytree(originalRoot/'workspace/import-sources/universities/stage-1',checker.ROOT/'workspace/import-sources/universities/stage-1');shutil.copytree(originalReport,checker.REPORT)
  # This checker covers the immutable research batch, before authorized exclusions.
  for artifact in json.loads((checker.REPORT/'summary.json').read_text())['universityFiles']:
   (checker.ROOT/artifact['path']).write_bytes(subprocess.check_output(['git','show','8a1bb4e6:'+artifact['path']],cwd=originalRoot))
  for path in subprocess.check_output(['git','ls-tree','-r','--name-only','8a1bb4e6','--','workspace/reference-data/cities','workspace/reference-data/regions','workspace/reference-data/countries'],cwd=originalRoot,text=True).splitlines():
   (checker.ROOT/path).write_bytes(subprocess.check_output(['git','show','8a1bb4e6:'+path],cwd=originalRoot))
 def tearDown(self):checker.ROOT,checker.REPORT=originalRoot,originalReport;self.tmp.cleanup()
 def updateHash(self,path,kind):
  summary=json.loads((checker.REPORT/'summary.json').read_text())
  for f in summary[kind]:
   if f['path']==str(path.relative_to(checker.ROOT)):f['afterSha256']=hashlib.sha256(path.read_bytes()).hexdigest()
  (checker.REPORT/'summary.json').write_text(json.dumps(summary))
 def editWorkbook(self,edit):
  summary=json.loads((checker.REPORT/'summary.json').read_text());p=checker.ROOT/summary['universityFiles'][0]['path'];w=openpyxl.load_workbook(p);s=w.worksheets[0];hi=next(i for i,r in enumerate(s.values,1) if 'Reference ID' in r);headers=[c.value for c in s[hi]];edit(s,hi,headers);w.save(p);self.updateHash(p,'universityFiles')
 def test_foreign_city_binding(self):
  def edit(s,hi,headers):
   row=hi+1;country=s.cell(row,headers.index('ISO3')+1).value
   city=next(r for p in (checker.ROOT/'workspace/reference-data/cities').glob('*/*.csv') for r in checker.csv_rows(p.read_bytes()) if r['countryIso3']!=country)
   s.cell(row,headers.index('City Source ID')+1,city['cityId'])
  self.editWorkbook(edit)
  with self.assertRaisesRegex(ValueError,'University city/country/region disagreement'):checker.verify()
 def test_deleting_an_institution_to_force_matching_is_detected(self):
  self.editWorkbook(lambda s,hi,headers:s.delete_rows(hi+1))
  with self.assertRaisesRegex(ValueError,'University identity loss'):checker.verify()
 def test_address_only_locality_must_not_invent_coordinates(self):
  additions=checker.records('city-additions.jsonl');a=next(r for r in additions if r['geography']['geo']['feature']=='OFFICIAL_ADDRESS_LOCALITY');p=checker.ROOT/a['file'];rows=checker.csv_rows(p.read_bytes());row=next(r for r in rows if r['cityId']==a['cityId']);row['latitude']='1';a['row']=dict(row);a['geography']['geo']['lat']='1';buf=io.StringIO();writer=csv.DictWriter(buf,fieldnames=list(rows[0]),lineterminator='\n');writer.writeheader();writer.writerows(rows);p.write_text(buf.getvalue());self.updateHash(p,'cityFiles');(checker.REPORT/'city-additions.jsonl').write_text(''.join(json.dumps(r)+'\n' for r in additions))
  with self.assertRaisesRegex(ValueError,'Official address must not invent coordinates'):checker.verify()
if __name__=='__main__':unittest.main()

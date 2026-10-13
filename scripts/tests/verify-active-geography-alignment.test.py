import hashlib
import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path
import openpyxl

script = Path(__file__).resolve().parents[1] / 'verify-active-geography-alignment.py'
spec = importlib.util.spec_from_file_location('alignment', script)
checker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(checker)
original_root, original_report = checker.ROOT, checker.REPORT

class AlignmentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        checker.ROOT = Path(self.tmp.name)
        checker.REPORT = checker.ROOT / original_report.relative_to(original_root)
        (checker.ROOT / '.git').symlink_to(original_root / '.git', target_is_directory=True)
        for directory in ('workspace/reference-data', 'workspace/import-sources/universities/stage-1', str(original_report.relative_to(original_root))):
            shutil.copytree(original_root / directory, checker.ROOT / directory)

    def tearDown(self):
        checker.ROOT, checker.REPORT = original_root, original_report
        self.tmp.cleanup()

    def edit(self, archived, mutation):
        summary = json.loads((checker.REPORT / 'summary.json').read_text())
        artifact = summary['universityFiles'][0]
        path = checker.ROOT / artifact['archivePath' if archived else 'path']
        workbook = openpyxl.load_workbook(path)
        sheet = workbook.worksheets[0]
        header = next(i for i, values in enumerate(sheet.values, 1) if 'Reference ID' in values)
        mutation(sheet, header, [cell.value for cell in sheet[header]])
        workbook.save(path)
        workbook.close()
        artifact['archiveSha256' if archived else 'afterSha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
        (checker.REPORT / 'summary.json').write_text(json.dumps(summary))

    def test_deleting_accepted_university_is_detected(self):
        self.edit(False, lambda sheet, header, columns: sheet.delete_rows(header + 1))
        with self.assertRaisesRegex(ValueError, 'University identity conservation'):
            checker.verify()

    def test_changing_city_label_is_detected(self):
        self.edit(False, lambda sheet, header, columns: setattr(sheet.cell(header + 1, columns.index('City') + 1), 'value', 'Wrong city'))
        with self.assertRaisesRegex(ValueError, 'Unexpected active university edit'):
            checker.verify()

    def test_archive_original_fields_are_preserved(self):
        self.edit(True, lambda sheet, header, columns: setattr(sheet.cell(header + 1, columns.index('Official English Name') + 1), 'value', 'Changed name'))
        with self.assertRaisesRegex(ValueError, 'Archive data loss'):
            checker.verify()

if __name__ == '__main__':
    unittest.main()

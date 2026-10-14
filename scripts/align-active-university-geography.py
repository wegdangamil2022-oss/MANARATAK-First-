#!/usr/bin/env python3
"""Exclude unresolved source rows, preserving complete records outside import inputs."""
import hashlib
import json
import subprocess
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'workspace/reports/active-geography-alignment/2026-10-13'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def apply():
    if REPORT.exists():
        raise ValueError('Alignment report already exists; refusing a second mutation')
    baseline = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
    REPORT.mkdir(parents=True)
    receipts, artifacts = [], []
    for path in sorted((ROOT / 'workspace/import-sources/universities/stage-1').glob('*.xlsx')):
        before_hash = sha(path)
        workbook = openpyxl.load_workbook(path)
        review = workbook['Geography Review']
        header_index = next(i for i, row in enumerate(review.values, 1) if 'Reference ID' in row)
        headers = [cell.value for cell in review[header_index]]
        archive = openpyxl.Workbook()
        sheet = archive.active
        sheet.title = 'Geography Review'
        for values in review.values:
            sheet.append(list(values))
        count = 0
        for values in list(review.values)[header_index:]:
            row = dict(zip(headers, values))
            if not row.get('Reference ID'):
                continue
            receipts.append({'referenceId': row['Reference ID'], 'sourcePath': str(path.relative_to(ROOT)),
                             'action': 'EXCLUDE_FROM_ACTIVE_IMPORT', 'reason': row['Geography Review Reason'],
                             'originalRow': row})
            count += 1
        archive_path = REPORT / path.name
        archive.save(archive_path)
        archive.close()
        review.delete_rows(header_index + 1, review.max_row - header_index)
        workbook.save(path)
        workbook.close()
        artifacts.append({'path': str(path.relative_to(ROOT)), 'beforeSha256': before_hash,
                          'afterSha256': sha(path), 'archivePath': str(archive_path.relative_to(ROOT)),
                          'archiveSha256': sha(archive_path), 'excludedRows': count})
    (REPORT / 'excluded-universities.jsonl').write_text(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in receipts))
    (REPORT / 'summary.json').write_text(json.dumps({'baseCommit': baseline,
        'authorization': 'User authorized correcting, adding or deleting mismatched source data on 2026-10-13.',
        'scope': 'ACTIVE_SOURCE_FILES_ONLY', 'excludedUniversities': len(receipts),
        'universityFiles': artifacts, 'databaseWrites': 0}, indent=2) + '\n')
    print(json.dumps({'excludedUniversities': len(receipts), 'databaseWrites': 0}))

if __name__ == '__main__':
    apply()
